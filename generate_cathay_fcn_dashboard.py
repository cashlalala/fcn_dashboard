from __future__ import annotations

import argparse
import html
import json
import re
from pathlib import Path
from typing import Iterable

import pandas as pd
import pdfplumber


DEFAULT_INPUT_DIR = Path(r"C:\Users\cash\OneDrive\國泰fcn報價")
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent / "output"

# Short code (used in output file names and --product) -> product type and display label.
PRODUCTS = {
    "fcn": {"product_type": "固定配息率觸及型", "label": "FCN"},
    "eln": {"product_type": "上限型", "label": "ELN"},
}


class NoQuotesError(ValueError):
    """The chosen stock has no quotes for the chosen product."""


def product_label(product_type: str) -> str:
    for product in PRODUCTS.values():
        if product["product_type"] == product_type:
            return product["label"]
    return product_type


DATE_RE = re.compile(r"(\d{4})[/.](\d{1,2})[/.](\d{1,2})")
PERIOD_RE = re.compile(
    r"有效期間[：:]\s*(\d{4}[/.]\d{1,2}[/.]\d{1,2})\s*~\s*(\d{4}[/.]\d{1,2}[/.]\d{1,2})"
)
QUOTE_LINE_RE = re.compile(
    r"^(?P<trading_identity>限PI承作|-)\s+"
    r"(?P<product_type>\S+)\s+"
    r"(?P<stock_id>\d{4})\s+"
    r"(?P<stock_name>\S+)\s+"
    r"(?P<tenor_days>\d+)\s+"
    r"(?P<lower_bound>[0-9.]+%?)\s+"
    r"(?P<upper_bound>[0-9.]+%?)\s+"
    r"(?P<coupon_rate>[0-9.]+%?)\s+"
    r"(?P<issue_price>[0-9.]+%?)\s+"
    r"(?P<participation_rate>-|[0-9.]+%?)\s+"
    r"(?P<ko_start>-|\d+)\s+"
    r"(?P<coupon_period>-|\d+)"
)


def normalize_date(value: str) -> str:
    match = DATE_RE.search(value)
    if not match:
        raise ValueError(f"Cannot parse date from: {value}")
    year, month, day = match.groups()
    return f"{int(year):04d}-{int(month):02d}-{int(day):02d}"


def parse_effective_period(text: str) -> tuple[str | None, str | None]:
    period_match = PERIOD_RE.search(text)
    if period_match:
        return normalize_date(period_match.group(1)), normalize_date(period_match.group(2))

    if "有效至" in text:
        date_match = DATE_RE.search(text)
        if date_match:
            return None, normalize_date(date_match.group(0))

    return None, None


def parse_pct(value: str) -> float | None:
    if value == "-":
        return None
    return float(value.rstrip("%"))


def parse_int(value: str) -> int | None:
    if value == "-":
        return None
    return int(value)


def parse_quote_line(
    line: str,
    quote_start_date: str | None,
    quote_end_date: str | None,
    source_file: str,
) -> dict | None:
    match = QUOTE_LINE_RE.match(line.strip())
    if not match:
        return None

    values = match.groupdict()
    return {
        "source_file": source_file,
        "quote_start_date": quote_start_date or quote_end_date,
        "quote_end_date": quote_end_date,
        "trading_identity": values["trading_identity"],
        "product_type": values["product_type"],
        "stock_id": values["stock_id"],
        "stock_name": values["stock_name"],
        "tenor_days": int(values["tenor_days"]),
        "lower_bound_pct": parse_pct(values["lower_bound"]),
        "upper_bound_pct": parse_pct(values["upper_bound"]),
        "coupon_rate_pct": parse_pct(values["coupon_rate"]),
        "issue_price_pct": parse_pct(values["issue_price"]),
        "participation_rate": parse_pct(values["participation_rate"]),
        "ko_start_days": parse_int(values["ko_start"]),
        "coupon_period_days": parse_int(values["coupon_period"]),
    }


def extract_rows_from_pdf(pdf_path: Path) -> list[dict]:
    with pdfplumber.open(str(pdf_path)) as pdf:
        text = "\n".join(
            page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            for page in pdf.pages
        )

    quote_start_date, quote_end_date = parse_effective_period(text)
    rows = []
    for line in text.splitlines():
        row = parse_quote_line(line, quote_start_date, quote_end_date, pdf_path.name)
        if row:
            rows.append(row)
    return rows


def extract_rows(input_dir: Path) -> list[dict]:
    rows: list[dict] = []
    for pdf_path in sorted(input_dir.glob("*.pdf")):
        rows.extend(extract_rows_from_pdf(pdf_path))
    return rows


def build_dashboard_model(
    rows: Iterable[dict],
    stock_id: str = "2330",
    product_type: str = "固定配息率觸及型",
) -> dict:
    df = pd.DataFrame(list(rows))
    if df.empty:
        return {
            "stock_id": stock_id,
            "stock_name": "",
            "product_type": product_type,
            "date_labels": [],
            "series": [],
        }

    filtered = df.loc[
        df["stock_id"].astype(str).eq(str(stock_id)) & df["product_type"].eq(product_type)
    ].copy()
    if filtered.empty:
        return {
            "stock_id": stock_id,
            "stock_name": "",
            "product_type": product_type,
            "date_labels": [],
            "series": [],
        }

    filtered["quote_start_date"] = filtered["quote_start_date"].fillna(filtered["quote_end_date"])
    filtered = filtered.dropna(
        subset=["quote_start_date", "lower_bound_pct", "upper_bound_pct", "coupon_rate_pct"]
    )

    date_labels = sorted(filtered["quote_start_date"].unique().tolist())
    stock_name = str(filtered["stock_name"].dropna().iloc[0])

    grouped = (
        filtered.groupby(["lower_bound_pct", "upper_bound_pct", "quote_start_date"], as_index=False)
        .agg({"coupon_rate_pct": "max"})
        .sort_values(["lower_bound_pct", "upper_bound_pct", "quote_start_date"])
    )

    series = []
    for (lower, upper), group in grouped.groupby(["lower_bound_pct", "upper_bound_pct"]):
        points = [
            [str(row.quote_start_date), float(row.coupon_rate_pct)]
            for row in group.itertuples(index=False)
        ]
        series.append(
            {
                "label": f"{format_bound(lower)} / {format_bound(upper)}",
                "lower_bound_pct": float(lower),
                "upper_bound_pct": float(upper),
                "points": points,
            }
        )

    return {
        "stock_id": str(stock_id),
        "stock_name": stock_name,
        "product_type": product_type,
        "date_labels": date_labels,
        "series": series,
    }


def format_bound(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else f"{value:.2f}".rstrip("0").rstrip(".")


def build_pivot(rows: Iterable[dict], stock_id: str = "2330", product_type: str = "固定配息率觸及型") -> pd.DataFrame:
    model = build_dashboard_model(rows, stock_id=stock_id, product_type=product_type)
    records = []
    for item in model["series"]:
        record = {
            "履約價(%)": item["lower_bound_pct"],
            "上限價(%)": item["upper_bound_pct"],
        }
        for date, rate in item["points"]:
            record[date] = rate
        records.append(record)
    columns = ["履約價(%)", "上限價(%)"] + model["date_labels"]
    return pd.DataFrame(records, columns=columns)


def write_excel(rows: list[dict], output_path: Path, stock_id: str = "2330") -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    all_rows = pd.DataFrame(rows)
    target_rows = all_rows.loc[all_rows["stock_id"].astype(str).eq(str(stock_id))].copy()
    fcn_pivot = build_pivot(rows, stock_id=stock_id, product_type="固定配息率觸及型")
    eln_pivot = build_pivot(rows, stock_id=stock_id, product_type="上限型")

    summary = pd.DataFrame(
        [
            summarize_sheet("FCN 固定配息率觸及型", fcn_pivot),
            summarize_sheet("ELN 上限型", eln_pivot),
        ]
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False)
        fcn_pivot.to_excel(writer, sheet_name=f"{stock_id}_FCN_利率變化", index=False)
        eln_pivot.to_excel(writer, sheet_name=f"{stock_id}_ELN_利率變化", index=False)
        target_rows.to_excel(writer, sheet_name=f"{stock_id}_明細", index=False)

    format_excel(output_path)


def summarize_sheet(label: str, df: pd.DataFrame) -> dict:
    date_cols = [col for col in df.columns if re.match(r"\d{4}-\d{2}-\d{2}", str(col))]
    return {
        "商品": label,
        "上下限組合數": len(df),
        "報價日期數": len(date_cols),
        "最早日期": min(date_cols) if date_cols else None,
        "最新日期": max(date_cols) if date_cols else None,
    }


def format_excel(output_path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = load_workbook(output_path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "C2" if ws.max_column > 3 else "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for column_cells in ws.columns:
            values = [str(cell.value) for cell in column_cells if cell.value is not None]
            width = min(max((len(value) for value in values), default=8) + 2, 18)
            ws.column_dimensions[get_column_letter(column_cells[0].column)].width = width

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"

    wb.save(output_path)


def render_html(model: dict) -> str:
    safe_json = html.escape(json.dumps(model, ensure_ascii=False), quote=False)
    label = product_label(model["product_type"])
    title = f"{model['stock_id']} {model['stock_name']} {label} 利率趨勢"
    return f"""<!doctype html>
<html lang="zh-Hant">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{html.escape(title)}</title>
  <style>
    :root {{
      color-scheme: light;
      --ink: #172033;
      --muted: #667085;
      --line: #d9e1ea;
      --panel: #f7f9fc;
      --accent: #0f766e;
      --accent-soft: #d9f2ed;
      --danger: #b42318;
    }}
    * {{ box-sizing: border-box; }}
    body {{
      margin: 0;
      font-family: "Microsoft JhengHei", "Noto Sans TC", Arial, sans-serif;
      color: var(--ink);
      background: #ffffff;
    }}
    header {{
      padding: 22px 28px 16px;
      border-bottom: 1px solid var(--line);
      background: #ffffff;
    }}
    h1 {{
      margin: 0 0 8px;
      font-size: 24px;
      line-height: 1.25;
      letter-spacing: 0;
    }}
    .meta {{
      display: flex;
      gap: 12px;
      flex-wrap: wrap;
      color: var(--muted);
      font-size: 13px;
    }}
    main {{
      display: grid;
      grid-template-columns: minmax(260px, 340px) 1fr;
      min-height: calc(100vh - 88px);
    }}
    aside {{
      border-right: 1px solid var(--line);
      background: var(--panel);
      padding: 16px;
      overflow: auto;
    }}
    .tools {{
      display: grid;
      gap: 10px;
      margin-bottom: 12px;
    }}
    input {{
      width: 100%;
      height: 36px;
      border: 1px solid var(--line);
      border-radius: 6px;
      padding: 0 10px;
      font-size: 14px;
      background: #ffffff;
    }}
    .series-list {{
      display: grid;
      gap: 6px;
    }}
    button.series-button {{
      display: grid;
      grid-template-columns: 1fr auto;
      gap: 10px;
      width: 100%;
      min-height: 42px;
      border: 1px solid transparent;
      border-radius: 6px;
      padding: 8px 10px;
      background: #ffffff;
      color: var(--ink);
      cursor: pointer;
      text-align: left;
      font-size: 14px;
    }}
    button.series-button:hover {{
      border-color: #a5b4c3;
    }}
    button.series-button.active {{
      border-color: var(--accent);
      background: var(--accent-soft);
    }}
    .points-count {{
      color: var(--muted);
      font-size: 12px;
      align-self: center;
    }}
    section.chart-area {{
      padding: 22px 28px;
      min-width: 0;
    }}
    .chart-head {{
      display: flex;
      justify-content: space-between;
      gap: 16px;
      align-items: flex-end;
      margin-bottom: 14px;
    }}
    h2 {{
      margin: 0;
      font-size: 20px;
      line-height: 1.25;
      letter-spacing: 0;
    }}
    .stats {{
      color: var(--muted);
      font-size: 13px;
      text-align: right;
    }}
    .chart-wrap {{
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 14px;
      overflow: hidden;
      background: #ffffff;
    }}
    svg {{
      display: block;
      width: 100%;
      height: 420px;
    }}
    .axis {{ stroke: #98a2b3; stroke-width: 1; }}
    .grid {{ stroke: #edf1f5; stroke-width: 1; }}
    .trend {{ fill: none; stroke: var(--accent); stroke-width: 3; }}
    .dot {{ fill: #ffffff; stroke: var(--accent); stroke-width: 2; }}
    .axis-label, .dot-label {{
      fill: var(--muted);
      font-size: 12px;
    }}
    .rate-label {{
      fill: var(--ink);
      font-size: 12px;
      font-weight: 700;
    }}
    table {{
      width: 100%;
      border-collapse: collapse;
      margin-top: 18px;
      font-size: 14px;
    }}
    th, td {{
      border-bottom: 1px solid var(--line);
      padding: 9px 8px;
      text-align: right;
    }}
    th:first-child, td:first-child {{ text-align: left; }}
    th {{
      color: var(--muted);
      font-weight: 700;
      background: #fbfcfe;
    }}
    .empty {{
      padding: 32px;
      color: var(--danger);
      border: 1px solid var(--line);
      border-radius: 8px;
      background: #fff7f6;
    }}
    @media (max-width: 860px) {{
      main {{ grid-template-columns: 1fr; }}
      aside {{ border-right: 0; border-bottom: 1px solid var(--line); max-height: 42vh; }}
      section.chart-area {{ padding: 18px; }}
      .chart-head {{ align-items: flex-start; flex-direction: column; }}
      .stats {{ text-align: left; }}
      svg {{ height: 340px; }}
    }}
  </style>
</head>
<body>
  <header>
    <h1>{html.escape(title)}</h1>
    <div class="meta">
      <span>商品：{html.escape(model['product_type'])}</span>
      <span>報價日期：{len(model['date_labels'])}</span>
      <span>上下限組合：{len(model['series'])}</span>
    </div>
  </header>
  <main>
    <aside>
      <div class="tools">
        <input id="filter" type="search" placeholder="搜尋上下限，例如 95 / 105" aria-label="搜尋上下限">
      </div>
      <div id="series-list" class="series-list"></div>
    </aside>
    <section class="chart-area">
      <div class="chart-head">
        <div>
          <h2 id="chart-title">選擇一組上下限</h2>
          <div class="meta" id="chart-subtitle"></div>
        </div>
        <div class="stats" id="chart-stats"></div>
      </div>
      <div class="chart-wrap">
        <svg id="chart" viewBox="0 0 900 420" role="img" aria-label="利率趨勢圖"></svg>
      </div>
      <table id="point-table">
        <thead><tr><th>報價日期</th><th>年化配息率(%)</th></tr></thead>
        <tbody></tbody>
      </table>
    </section>
  </main>
  <script id="dashboard-data" type="application/json">{safe_json}</script>
  <script>
    const model = JSON.parse(document.getElementById("dashboard-data").textContent);
    const list = document.getElementById("series-list");
    const filter = document.getElementById("filter");
    const svg = document.getElementById("chart");
    const tableBody = document.querySelector("#point-table tbody");
    const title = document.getElementById("chart-title");
    const subtitle = document.getElementById("chart-subtitle");
    const stats = document.getElementById("chart-stats");
    let activeIndex = 0;

    function createButton(item, index) {{
      const button = document.createElement("button");
      button.type = "button";
      button.className = "series-button";
      button.setAttribute("data-series-index", index);
      button.innerHTML = `<span>${{item.label}}</span><span class="points-count">${{item.points.length}} 點</span>`;
      button.addEventListener("click", () => renderSeries(index));
      return button;
    }}

    function normalizeSearch(text) {{
      // Ignore "/" and spacing so "95/105", "95 105" and "95 / 105" all match anywhere in the label.
      return text.split("/").join(" ").split(" ").filter(Boolean).join(" ");
    }}

    function renderList() {{
      const keyword = normalizeSearch(filter.value);
      list.innerHTML = "";
      model.series.forEach((item, index) => {{
        if (keyword && !normalizeSearch(item.label).includes(keyword)) return;
        list.appendChild(createButton(item, index));
      }});
      markActive();
    }}

    function markActive() {{
      document.querySelectorAll(".series-button").forEach((button) => {{
        button.classList.toggle("active", Number(button.dataset.seriesIndex) === activeIndex);
      }});
    }}

    function renderSeries(index) {{
      activeIndex = index;
      markActive();
      const item = model.series[index];
      title.textContent = `履約價 / 上限價：${{item.label}}`;
      subtitle.textContent = `${{model.stock_id}} ${{model.stock_name}}｜${{model.product_type}}`;
      if (!item || item.points.length === 0) {{
        svg.innerHTML = `<text x="30" y="50" class="axis-label">此上下限組合沒有可繪製資料</text>`;
        tableBody.innerHTML = "";
        stats.textContent = "";
        return;
      }}

      const rates = item.points.map((point) => point[1]);
      const minRate = Math.min(...rates);
      const maxRate = Math.max(...rates);
      const latest = item.points[item.points.length - 1];
      stats.textContent = `最新：${{latest[0]}}｜${{latest[1].toFixed(2)}}%｜區間：${{minRate.toFixed(2)}}% - ${{maxRate.toFixed(2)}}%`;
      drawChart(item.points, minRate, maxRate);
      renderTable(item.points);
    }}

    function drawChart(points, minRate, maxRate) {{
      const width = 900;
      const height = 420;
      const margin = {{ left: 64, right: 24, top: 28, bottom: 72 }};
      const innerW = width - margin.left - margin.right;
      const innerH = height - margin.top - margin.bottom;
      const span = Math.max(maxRate - minRate, 1);
      const yMin = Math.max(0, minRate - span * 0.18);
      const yMax = maxRate + span * 0.18;
      const ySpan = Math.max(yMax - yMin, 1);
      const x = (i) => margin.left + (points.length === 1 ? innerW / 2 : (i / (points.length - 1)) * innerW);
      const y = (rate) => margin.top + (1 - (rate - yMin) / ySpan) * innerH;

      const path = points.map((point, i) => `${{i === 0 ? "M" : "L"}} ${{x(i).toFixed(1)}} ${{y(point[1]).toFixed(1)}}`).join(" ");
      const grid = [0, 0.25, 0.5, 0.75, 1].map((ratio) => {{
        const yy = margin.top + ratio * innerH;
        const label = yMax - ratio * ySpan;
        return `<line x1="${{margin.left}}" y1="${{yy}}" x2="${{width - margin.right}}" y2="${{yy}}" class="grid"></line>
          <text x="${{margin.left - 10}}" y="${{yy + 4}}" text-anchor="end" class="axis-label">${{label.toFixed(1)}}%</text>`;
      }}).join("");
      const dots = points.map((point, i) => {{
        const xx = x(i);
        const yy = y(point[1]);
        return `<circle cx="${{xx}}" cy="${{yy}}" r="4.5" class="dot"></circle>
          <text x="${{xx}}" y="${{yy - 10}}" text-anchor="middle" class="rate-label">${{point[1].toFixed(2)}}%</text>
          <text x="${{xx}}" y="${{height - 28}}" text-anchor="middle" class="axis-label" transform="rotate(-35 ${{xx}} ${{height - 28}})">${{point[0].slice(5)}}</text>`;
      }}).join("");

      svg.innerHTML = `${{grid}}
        <line x1="${{margin.left}}" y1="${{height - margin.bottom}}" x2="${{width - margin.right}}" y2="${{height - margin.bottom}}" class="axis"></line>
        <line x1="${{margin.left}}" y1="${{margin.top}}" x2="${{margin.left}}" y2="${{height - margin.bottom}}" class="axis"></line>
        <path d="${{path}}" class="trend"></path>
        ${{dots}}`;
    }}

    function renderTable(points) {{
      tableBody.innerHTML = points.map((point) => `<tr><td>${{point[0]}}</td><td>${{point[1].toFixed(2)}}</td></tr>`).join("");
    }}

    filter.addEventListener("input", renderList);
    renderList();
    if (model.series.length) {{
      renderSeries(0);
    }} else {{
      document.querySelector(".chart-wrap").innerHTML = `<div class="empty">沒有找到 ${{model.stock_id}} {html.escape(label)} 報價資料。</div>`;
    }}
  </script>
</body>
</html>
"""


def write_outputs(
    input_dir: Path,
    output_dir: Path,
    stock_id: str = "2330",
    rows: list[dict] | None = None,
    product: str = "fcn",
) -> dict:
    if rows is None:
        rows = extract_rows(input_dir)
    product_type = PRODUCTS[product]["product_type"]
    model = build_dashboard_model(rows, stock_id=stock_id, product_type=product_type)
    if not model["series"]:
        raise NoQuotesError(f"股票 {stock_id} 沒有 {PRODUCTS[product]['label']}（{product_type}）報價。")
    output_dir.mkdir(parents=True, exist_ok=True)
    html_path = output_dir / f"{stock_id}_{product}_dashboard.html"
    excel_path = output_dir / f"cathay_fcn_{stock_id}_rates.xlsx"
    html_path.write_text(render_html(model), encoding="utf-8")
    write_excel(rows, excel_path, stock_id=stock_id)
    return {
        "rows": len(rows),
        "series": len(model["series"]),
        "dates": len(model["date_labels"]),
        "html": html_path,
        "excel": excel_path,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="產生國泰 FCN/ELN 報價互動 HTML 與 Excel。")
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--stock-id", default="2330")
    parser.add_argument("--product", choices=sorted(PRODUCTS), default="fcn", help="HTML 圖表的商品類型")
    args = parser.parse_args()

    try:
        result = write_outputs(args.input_dir, args.output_dir, stock_id=args.stock_id, product=args.product)
    except NoQuotesError as exc:
        raise SystemExit(str(exc)) from exc
    print(
        f"Wrote {result['html']} and {result['excel']} "
        f"from {result['rows']} rows; {result['series']} series across {result['dates']} dates."
    )


if __name__ == "__main__":
    main()
