import sys
import unittest
from pathlib import Path

PROJECT_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_DIR))

import tempfile  # noqa: E402

from generate_cathay_fcn_dashboard import (  # noqa: E402
    NoQuotesError,
    build_dashboard_model,
    parse_quote_line,
    render_html,
    write_outputs,
)

ELN_ROW = {"stock_id": "2345", "stock_name": "智邦", "product_type": "上限型", "quote_start_date": "2026-05-11", "quote_end_date": "2026-05-15", "lower_bound_pct": 100.0, "upper_bound_pct": 104.0, "coupon_rate_pct": 12.56}


class CathayFcnDashboardTests(unittest.TestCase):
    def test_parse_quote_line_normalizes_percent_values(self):
        row = parse_quote_line(
            "限PI承作 固定配息率觸及型 2330 台積電 84 91% 105% 10.96 100.00% 2 28 28",
            "2026-04-27",
            "2026-05-08",
            "sample.pdf",
        )

        self.assertEqual(row["stock_id"], "2330")
        self.assertEqual(row["stock_name"], "台積電")
        self.assertEqual(row["lower_bound_pct"], 91.0)
        self.assertEqual(row["upper_bound_pct"], 105.0)
        self.assertEqual(row["coupon_rate_pct"], 10.96)

    def test_build_dashboard_model_filters_2330_fcn_and_skips_missing_dates(self):
        rows = [
            {"stock_id": "2330", "stock_name": "台積電", "product_type": "固定配息率觸及型", "quote_start_date": "2026-04-27", "quote_end_date": "2026-05-08", "lower_bound_pct": 91.0, "upper_bound_pct": 105.0, "coupon_rate_pct": 10.96},
            {"stock_id": "2330", "stock_name": "台積電", "product_type": "固定配息率觸及型", "quote_start_date": "2026-05-25", "quote_end_date": "2026-05-29", "lower_bound_pct": 91.0, "upper_bound_pct": 105.0, "coupon_rate_pct": 13.53},
            {"stock_id": "2330", "stock_name": "台積電", "product_type": "固定配息率觸及型", "quote_start_date": "2026-05-25", "quote_end_date": "2026-05-29", "lower_bound_pct": 92.0, "upper_bound_pct": 106.0, "coupon_rate_pct": 16.63},
            {"stock_id": "2454", "stock_name": "聯發科", "product_type": "固定配息率觸及型", "quote_start_date": "2026-05-25", "quote_end_date": "2026-05-29", "lower_bound_pct": 91.0, "upper_bound_pct": 105.0, "coupon_rate_pct": 20.00},
        ]

        model = build_dashboard_model(rows)

        self.assertEqual(model["stock_id"], "2330")
        self.assertEqual(model["date_labels"], ["2026-04-27", "2026-05-25"])
        self.assertEqual(len(model["series"]), 2)
        first = model["series"][0]
        self.assertEqual(first["label"], "91 / 105")
        self.assertEqual(first["points"], [["2026-04-27", 10.96], ["2026-05-25", 13.53]])

    def test_render_html_embeds_data_and_clickable_controls(self):
        model = {
            "stock_id": "2330",
            "stock_name": "台積電",
            "product_type": "固定配息率觸及型",
            "date_labels": ["2026-04-27", "2026-05-25"],
            "series": [
                {
                    "label": "91 / 105",
                    "lower_bound_pct": 91.0,
                    "upper_bound_pct": 105.0,
                    "points": [["2026-04-27", 10.96], ["2026-05-25", 13.53]],
                }
            ],
        }

        html = render_html(model)

        self.assertIn('<script id="dashboard-data" type="application/json">', html)
        self.assertIn("renderSeries", html)
        self.assertIn("data-series-index", html)
        self.assertIn("91 / 105", html)

    def test_write_outputs_refuses_stock_without_quotes_for_product(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(NoQuotesError):
                write_outputs(Path(tmp), Path(tmp), stock_id="2345", rows=[ELN_ROW], product="fcn")
            self.assertEqual(list(Path(tmp).iterdir()), [])

    def test_write_outputs_eln_uses_stock_specific_names(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            result = write_outputs(Path(tmp), Path(tmp), stock_id="2345", rows=[ELN_ROW], product="eln")
            html = result["html"].read_text(encoding="utf-8")
            sheets = load_workbook(result["excel"]).sheetnames

        self.assertEqual(result["html"].name, "2345_eln_dashboard.html")
        self.assertIn("2345 智邦 ELN 利率趨勢", html)
        self.assertEqual(sheets, ["Summary", "2345_FCN_利率變化", "2345_ELN_利率變化", "2345_明細"])


if __name__ == "__main__":
    unittest.main()
